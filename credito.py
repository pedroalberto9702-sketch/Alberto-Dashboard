"""
Crédito disponível: lê a planilha do Tesouro no e-mail e publica os dados.

Roda dentro da rotina do GitHub Actions (mesma do coletor do Compras), depois
da coleta. Faz quatro coisas:

  1. entra na caixa de entrada por IMAP e procura o e-mail mais recente que
     tenha uma planilha .xlsx (filtrando por remetente e/ou assunto);
  2. lê a planilha, desfazendo as células mescladas do relatório do Tesouro;
  3. CIFRA o resultado com a senha CREDITO_SENHA (AES-256-GCM);
  4. grava dados/credito/credito.enc.json.

Por que cifrar: o repositório e o site são públicos, e esta planilha mostra o
orçamento por programa de trabalho, incluindo nomes de operações. Sem a
senha, o arquivo publicado é ilegível. A página pede a senha uma vez por
aparelho. Sem CREDITO_SENHA a rotina NÃO publica nada (comportamento seguro
por padrão); só publica em texto aberto se CREDITO_PUBLICO=sim for definido
de propósito.

Variáveis de ambiente (todas vêm dos Secrets do repositório):

  CREDITO_EMAIL        endereço da caixa de entrada (obrigatória)
  CREDITO_SENHA_APP    senha de aplicativo do Gmail (obrigatória)
  CREDITO_SENHA        senha que cifra o arquivo publicado (obrigatória)
  CREDITO_REMETENTE    filtra pelo remetente (opcional, trecho do endereço)
  CREDITO_ASSUNTO      filtra pelo assunto (opcional, trecho do texto)
  CREDITO_IMAP         servidor IMAP (padrão imap.gmail.com)
  CREDITO_DIAS         quantos dias para trás procurar (padrão 4)
  CREDITO_PUBLICO      "sim" para publicar sem cifrar (não recomendado)
  SAIDA                pasta de saída (padrão dados)

Código de saída: 0 se atualizou ou se não havia nada novo; 1 em erro de
verdade (login recusado, planilha com layout inesperado...).
"""

from __future__ import annotations

import base64
import datetime as dt
import email
import imaplib
import io
import json
import os
import sys
from email.header import decode_header, make_header
from email.utils import parsedate_to_datetime
from pathlib import Path

import openpyxl

COLUNAS = {
    "acao": "Ação Governo",
    "fonte": "Fonte Recursos Detalhada",
    "ptres": "PTRES",
    "pi": "PI",
}
ITERACOES = 200_000


# ---------------------------------------------------------------------------
# leitura da planilha
# ---------------------------------------------------------------------------

def _num(v) -> float:
    if v is None or v == "":
        return 0.0
    if isinstance(v, (int, float)):
        return round(float(v), 2)
    txt = str(v).strip().replace("R$", "").replace(" ", "")
    if "," in txt:
        txt = txt.replace(".", "").replace(",", ".")
    try:
        return round(float(txt), 2)
    except ValueError:
        return 0.0


def _txt(v) -> str:
    return "" if v is None else str(v).strip()


def ler_planilha(origem) -> dict:
    """Lê o relatório 'Crédito Disponível' do Tesouro Gerencial.

    `origem` é um caminho ou bytes. O relatório vem com células mescladas
    (Ação, Fonte e PTRES ocupam duas linhas quando repetem); aqui elas são
    desfeitas, para cada linha ficar completa.
    """
    wb = openpyxl.load_workbook(
        io.BytesIO(origem) if isinstance(origem, (bytes, bytearray)) else origem,
        data_only=True,
    )
    ws = wb.active

    preenchido = {}
    for faixa in ws.merged_cells.ranges:
        valor = ws.cell(faixa.min_row, faixa.min_col).value
        for r in range(faixa.min_row, faixa.max_row + 1):
            for c in range(faixa.min_col, faixa.max_col + 1):
                preenchido[(r, c)] = valor

    def celula(r, c):
        return preenchido.get((r, c), ws.cell(r, c).value)

    # cabeçalho: a linha onde está 'Ação Governo'
    cab = None
    for r in range(1, min(ws.max_row, 30) + 1):
        for c in range(1, ws.max_column + 1):
            if _txt(celula(r, c)) == COLUNAS["acao"]:
                cab = r
                break
        if cab:
            break
    if not cab:
        raise ValueError("Layout inesperado: não achei o cabeçalho 'Ação Governo'.")

    def coluna(nome, linhas=(0, 1)):
        for c in range(1, ws.max_column + 1):
            for d in linhas:
                if _txt(celula(cab + d, c)).lower() == nome.lower():
                    return c
        return None

    col = {k: coluna(v, (0,)) for k, v in COLUNAS.items()}
    col["nd"] = coluna("Natureza Despesa", (1, 0))
    col["recebido"] = coluna("Crédito Recebido", (0,))
    col["disponivel"] = coluna("CREDITO DISPONIVEL", (0,))
    col["empenhado"] = coluna("DESPESAS EMPENHADAS", (0,))
    faltam = [k for k, v in col.items() if v is None]
    if faltam:
        raise ValueError(f"Layout inesperado: colunas não encontradas: {faltam}")
    col["descricao"] = col["pi"] + 1      # a descrição do PI fica ao lado do código

    # UG: 'UG Responsável: NOME:120641' -> só o código
    ug = ""
    for r in range(1, cab):
        t = _txt(celula(r, 1))
        if t.lower().startswith("ug"):
            ug = t.rsplit(":", 1)[-1].strip()

    linhas = []
    for r in range(cab + 2, ws.max_row + 1):
        nd = _txt(celula(r, col["nd"]))
        recebido = celula(r, col["recebido"])
        if not nd and recebido in (None, ""):
            break
        linha = {
            "acao": _txt(celula(r, col["acao"])),
            "fonte": _txt(celula(r, col["fonte"])),
            "ptres": _txt(celula(r, col["ptres"])),
            "pi": _txt(celula(r, col["pi"])),
            "descricao": _txt(celula(r, col["descricao"])),
            "nd": nd,
            "recebido": _num(recebido),
            "disponivel": _num(celula(r, col["disponivel"])),
            "empenhado": _num(celula(r, col["empenhado"])),
        }
        linha["chave"] = "|".join(
            linha[k] for k in ("acao", "fonte", "ptres", "pi", "nd"))
        linhas.append(linha)

    if not linhas:
        raise ValueError("A planilha não tem nenhuma linha de crédito.")
    chaves = [l["chave"] for l in linhas]
    if len(set(chaves)) != len(chaves):
        raise ValueError("Há linhas repetidas (mesma ação, fonte, PTRES, PI e ND).")
    return {"ug": ug, "linhas": linhas}


# ---------------------------------------------------------------------------
# cifra (compatível com WebCrypto do navegador)
# ---------------------------------------------------------------------------

def cifrar(dados: dict, senha: str) -> dict:
    from cryptography.hazmat.primitives import hashes
    from cryptography.hazmat.primitives.ciphers.aead import AESGCM
    from cryptography.hazmat.primitives.kdf.pbkdf2 import PBKDF2HMAC

    sal, iv = os.urandom(16), os.urandom(12)
    chave = PBKDF2HMAC(
        algorithm=hashes.SHA256(), length=32, salt=sal, iterations=ITERACOES
    ).derive(senha.encode("utf-8"))
    bruto = json.dumps(dados, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
    ct = AESGCM(chave).encrypt(iv, bruto, None)       # ciphertext || tag
    b64 = lambda b: base64.b64encode(b).decode("ascii")
    return {"v": 1, "kdf": "PBKDF2-SHA256", "iter": ITERACOES,
            "salt": b64(sal), "iv": b64(iv), "ct": b64(ct)}


# ---------------------------------------------------------------------------
# e-mail
# ---------------------------------------------------------------------------

def _decodificar(valor) -> str:
    try:
        return str(make_header(decode_header(valor or "")))
    except Exception:
        return str(valor or "")


def _simples(t: str) -> str:
    """Minúsculas e sem acento, para comparar assunto e remetente."""
    import unicodedata
    t = unicodedata.normalize("NFKD", str(t or "")).encode("ascii", "ignore").decode()
    return t.lower()


def achar_planilha(host, usuario, senha_app, remetente, assunto, dias):
    """Devolve (bytes do .xlsx, data do e-mail, assunto) do mais recente.

    A busca no servidor é só por data. O filtro de remetente e assunto é
    feito aqui, sem acento: o IMAP trata mal texto acentuado na busca, e um
    assunto como 'Crédito Disponível' falharia.
    """
    imap = imaplib.IMAP4_SSL(host)
    imap.login(usuario, senha_app)
    try:
        imap.select("INBOX", readonly=True)
        desde = (dt.date.today() - dt.timedelta(days=dias)).strftime("%d-%b-%Y")
        tipo, achados = imap.search(None, "SINCE", desde)
        if tipo != "OK":
            raise RuntimeError("A busca no e-mail falhou.")
        for num in reversed(achados[0].split()):       # do mais novo para o mais velho
            tipo, cab = imap.fetch(num, "(BODY.PEEK[HEADER.FIELDS (FROM SUBJECT DATE)])")
            if tipo != "OK":
                continue
            h = email.message_from_bytes(cab[0][1])
            if remetente and _simples(remetente) not in _simples(_decodificar(h["From"])):
                continue
            if assunto and _simples(assunto) not in _simples(_decodificar(h["Subject"])):
                continue
            tipo, partes = imap.fetch(num, "(BODY.PEEK[])")
            if tipo != "OK":
                continue
            msg = email.message_from_bytes(partes[0][1])
            for parte in msg.walk():
                nome = _decodificar(parte.get_filename())
                if nome.lower().endswith(".xlsx"):
                    data = parsedate_to_datetime(msg["Date"]) if msg["Date"] else None
                    return (parte.get_payload(decode=True), data,
                            _decodificar(msg["Subject"]))
        return None
    finally:
        try:
            imap.logout()
        except Exception:
            pass


# ---------------------------------------------------------------------------

def main() -> int:
    env = os.environ.get
    usuario, senha_app = env("CREDITO_EMAIL"), env("CREDITO_SENHA_APP")
    if not usuario or not senha_app:
        print("crédito: sem credenciais de e-mail, etapa ignorada.")
        return 0

    senha = env("CREDITO_SENHA", "")
    publico = (env("CREDITO_PUBLICO", "").lower() == "sim")
    if not senha and not publico:
        print("crédito: defina o Secret CREDITO_SENHA para cifrar o arquivo "
              "(ou CREDITO_PUBLICO=sim para publicar em texto aberto). "
              "Nada foi publicado.", file=sys.stderr)
        return 1

    try:
        achado = achar_planilha(
            env("CREDITO_IMAP", "imap.gmail.com"), usuario, senha_app,
            env("CREDITO_REMETENTE", ""), env("CREDITO_ASSUNTO", ""),
            int(env("CREDITO_DIAS", "4")))
    except imaplib.IMAP4.error as e:
        print(f"crédito: o servidor de e-mail recusou o acesso ({e}).", file=sys.stderr)
        return 1

    if not achado:
        print("crédito: nenhum e-mail com planilha .xlsx no período; "
              "dados anteriores mantidos.")
        return 0

    bruto, data_email, assunto = achado
    try:
        dados = ler_planilha(bruto)
    except ValueError as e:
        print(f"crédito: {e}", file=sys.stderr)
        return 1

    dados["atualizado"] = (data_email or dt.datetime.now(dt.timezone.utc)).isoformat()
    dados["coletado"] = dt.datetime.now(dt.timezone.utc).isoformat()

    pasta = Path(env("SAIDA", "dados")) / "credito"
    pasta.mkdir(parents=True, exist_ok=True)
    cifrado, aberto = pasta / "credito.enc.json", pasta / "credito.json"
    if senha:
        cifrado.write_text(json.dumps(cifrar(dados, senha)), encoding="utf-8")
        aberto.unlink(missing_ok=True)       # nunca deixa o texto aberto ao lado
        print(f"crédito: {len(dados['linhas'])} linhas, e-mail de "
              f"{dados['atualizado']}; publicado cifrado.")
    else:
        aberto.write_text(json.dumps(dados, ensure_ascii=False), encoding="utf-8")
        cifrado.unlink(missing_ok=True)
        print(f"crédito: {len(dados['linhas'])} linhas publicadas EM TEXTO ABERTO.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
